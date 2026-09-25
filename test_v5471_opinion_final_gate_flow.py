from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app_core import build_gorus_docx
from gorus_audit import validate_gorus_template_fidelity, validate_opinion_narrative_rules
from rules import APP_VERSION, RULESET_VERSION, GORUS_RULES

ROOT = Path(__file__).resolve().parent


def _opinion():
    return {
        "application_no": "2024/005033",
        "applicant": "OSTİM TEKNİK ÜNİVERSİTESİ",
        "reference": "698947",
        "intro": "14.07.2026 tarihli Birinci Bildirimde, 1-3 numaralı istemlerin yenilik ve sanayiye uygulanabilirlik kriterlerini sağladığı, ancak buluş basamağı kriterini sağlamadığı belirtilmiştir. Başvuru sahibinin görüşleri aşağıda dikkatinize sunulmaktadır.",
        "cited_documents": [],
        "sections": [],
        "combined_assessment": {"heading": "", "paragraphs": []},
        "conclusion": ["Ayırt edici teknik fark, sensör verilerinin belirli bir işlevsel ilişki içinde değerlendirilmesidir. Teknik etki açıklanmıştır. Objektif teknik problem bu işlevsel ilişkinin nasıl sağlanacağıdır. Önceki teknikte bu yönde motivasyon veya yönlendirme bulunmamaktadır."],
        "signoff": "Saygılarımızla,\nDESTEK PATENT A.Ş.",
        "_output_language": "Türkçe",
        "_recipient_office": "TURKISH PATENT AND TRADEMARK OFFICE",
    }


def test_version_bumped():
    assert APP_VERSION == "v5.4.78"
    assert RULESET_VERSION == "2026-09-25.v70"


def test_binding_title_paragraphs_are_centered_and_gate_enforces_it():
    op = _opinion()
    data = build_gorus_docx(op)
    doc = Document(io.BytesIO(data))
    assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert doc.paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.CENTER
    validate_gorus_template_fidelity(data, ROOT / "Gorus_metni_696809_template.docx", op)

    doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    out = io.BytesIO(); doc.save(out)
    with pytest.raises(ValueError, match="hizalama"):
        validate_gorus_template_fidelity(out.getvalue(), ROOT / "Gorus_metni_696809_template.docx", op)


def test_internal_uppercase_how_label_is_rejected_but_natural_lowercase_is_allowed():
    op = _opinion()
    op["conclusion"] = ["Ayırt edici teknik farkın nasıl ortaya çıktığı unsur ilişkilerinden anlaşılmaktadır. Teknik etki vardır. Objektif teknik problem açıklanmıştır. Motivasyon ve yönlendirme bulunmamaktadır."]
    validate_opinion_narrative_rules(op, "buluş basamağı", "teknik katkı teknik etki objektif teknik problem motivasyon")
    op["conclusion"] = ["İstem 1'in elemanları arasındaki NASIL ilişkisi savunmayı desteklemektedir. Teknik etki vardır. Objektif teknik problem açıklanmıştır. Motivasyon ve yönlendirme bulunmamaktadır."]
    with pytest.raises(ValueError, match="iç-süreç etiketi"):
        validate_opinion_narrative_rules(op, "buluş basamağı", "teknik katkı teknik etki objektif teknik problem motivasyon")


def test_normal_flow_is_automatic_only_when_no_claim_amendment_is_required():
    app = (ROOT / "app.py").read_text(encoding="utf-8")
    assert 'revision_status = "Mevcut istem seti üzerinden otomatik revizyonsuz görüş"' in app
    assert 'if not (st.session_state.gorus_opinion_data and st.session_state.gorus_opinion_status == revision_status):' in app
    assert '"Önerilen istem revizyonunu uygula"' in app
    assert '"Mevcut istemlerle devam et"' in app
    assert 'ready_to_generate = False' in app
    assert "### Görüşü revize et" in app


def test_rules_require_amendment_decision_gate_and_internal_label_gate():
    assert "GÖRÜŞ AKIŞ / İSTEM REVİZYONU KARAR KAPISI" in GORUS_RULES
    assert "NİHAİ GÖRÜŞ İÇ-SÜREÇ ETİKETİ KAPISI" in GORUS_RULES
    assert "GÖRÜŞ KURUM BAŞLIĞI HİZALAMA KAPISI" in GORUS_RULES


def _utility_model_opinion():
    return {
        "application_no": "2024/018834",
        "applicant": "ASELSAN ELEKTRONİK SANAYİ VE TİCARET ANONİM ŞİRKETİ",
        "reference": "700286",
        "intro": "20.01.2026 tarihli araştırma raporunda, 1 ve 2 numaralı istemlerin yenilik kriterini sağlamadığı belirtilmiştir. Başvuru sahibinin görüşleri aşağıda dikkatinize sunulmaktadır. Araştırma raporunda 1 ve 2 numaralı istemler bakımından gösterilen benzer dokümanlar aşağıdadır: D1",
        "amendment_assessment": {"heading": "", "blocks": []},
        "cited_documents": [{"label": "D1", "number": "CN117039321A", "category": "X"}],
        "sections": [{
            "label": "D1",
            "heading": "D1 (CN117039321A) dokümanı:",
            "blocks": [{"type": "paragraph", "text": "D1 bir batarya kapağı düzeneğini açıklamaktadır."}],
            "novelty_heading": "",
            "novelty_paragraphs": ["Ayırt edici teknik fark, istemdeki yaylı pin ile kapak içindeki iletken plaka arasındaki doğrudan temas ilişkisidir. D1 bu yapısal ilişkiyi doğrudan ve açık biçimde açıklamamaktadır."],
            "inventive_step_heading": "",
            "inventive_step_paragraphs": [],
        }],
        "combined_assessment": {"heading": "", "paragraphs": [], "groups": []},
        "conclusion": ["Bu nedenle mevcut istemlerin D1 karşısında yenilik yönünden yeniden değerlendirilmesi gerektiği düşünülmektedir."],
        "signoff": "Saygılarımızla,\nDESTEK PATENT A.Ş.",
    }


def test_utility_model_opinion_is_novelty_only_and_rejects_ai_chain_symbols():
    from gorus_audit import is_utility_model_search_report
    report = "TÜRK PATENT VE MARKA KURUMU\nFAYDALI MODEL ARAŞTIRMA RAPORU\nD1: CN117039321A"
    assert is_utility_model_search_report(report)
    op = _utility_model_opinion()
    validate_opinion_narrative_rules(op, report, "ayırt edici teknik fark")
    op["sections"][0]["novelty_paragraphs"] = ["Yaylı pin → vida → plaka biçiminde bir zincir olduğu ileri sürülmüştür."]
    with pytest.raises(ValueError, match="ok işaretli teknik zincir"):
        validate_opinion_narrative_rules(op, report, "ayırt edici teknik fark")


def test_utility_model_opinion_rejects_inventive_step_and_claim_transfer_meta_language():
    report = "TÜRK PATENT VE MARKA KURUMU\nFAYDALI MODEL ARAŞTIRMA RAPORU\nD1: CN117039321A"
    op = _utility_model_opinion()
    op["sections"][0]["inventive_step_paragraphs"] = ["Buluş basamağı bakımından ayrıca değerlendirme yapılmalıdır."]
    with pytest.raises(ValueError, match="Faydalı model görüş kapısı"):
        validate_opinion_narrative_rules(op, report, "ayırt edici teknik fark")
    op = _utility_model_opinion()
    op["conclusion"] = ["İstem 3 özelliğinin ana isteme taşınması gerekmez. Ayırt edici teknik fark korunmaktadır."]
    with pytest.raises(ValueError, match="Revizyonsuz görüş kapısı"):
        validate_opinion_narrative_rules(op, report, "ayırt edici teknik fark")


def test_rules_include_utility_model_novelty_only_and_original_figure_extraction():
    assert "FAYDALI MODEL ARAŞTIRMA RAPORU ÖZEL KAPISI" in GORUS_RULES
    assert "NİHAİ GÖRÜŞ DOĞAL PATENT DİLİ KAPISI" in GORUS_RULES
    assert "REVİZYONSUZ GÖRÜŞTE İSTEM-TAŞIMA META DİLİ YASAĞI" in GORUS_RULES
    assert "D-DOKÜMANI ŞEKİL KULLANIM KAPISI" in GORUS_RULES
