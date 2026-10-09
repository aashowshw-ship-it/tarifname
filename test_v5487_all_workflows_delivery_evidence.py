from __future__ import annotations
import io
import pytest
from docx import Document
from rules import FINAL_COMPLIANCE_REQUIRED_CHECKS, final_compliance_gate
from release_evidence import issue_delivery_evidence


def tiny_docx(kind):
    d=Document()
    text={'gorus':'GÖRÜŞ D1 D2','tip3':'ÖN ARAŞTIRMA RAPORU','tip3_update':'ÖN ARAŞTIRMA RAPORU',
          'claim_amendment':'Düzenlenmiş patent istemi', 'tarifname_update':'Güncellenen patent açıklaması',
          'figures':'ŞEKİL 1', 'figure_update':'ŞEKİL 1'}.get(kind,'Metin')
    d.add_paragraph(text)
    if kind in ('figures','figure_update'):
        from docx.oxml import OxmlElement
        drawing=OxmlElement('w:drawing'); d.paragraphs[0]._p.append(drawing)
    if kind=='tarifname_update':
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        ins=OxmlElement('w:ins');ins.set(qn('w:author'),'Destek Patent')
        d.paragraphs[0]._p.append(ins)
    b=io.BytesIO();d.save(b);return b.getvalue()


def full_checks(kind):
    return {k:True for k in FINAL_COMPLIANCE_REQUIRED_CHECKS[kind]}


def test_existing_bool_receipts_never_bypass_gate():
    for kind in FINAL_COMPLIANCE_REQUIRED_CHECKS:
        data=tiny_docx(kind if kind!='tarifname' else 'tip3')
        with pytest.raises(ValueError, match='makbuzu yok'):
            final_compliance_gate(kind,data=data,output_name='test.docx',default_name='test.docx',checks=full_checks(kind))


def test_evidence_is_bound_to_actual_word_and_filename():
    kind='gorus';original=tiny_docx(kind);checks=full_checks(kind)
    proof=issue_delivery_evidence(kind,original,[b'uploaded official examination report'],checks,'Görüş Metni_777.docx')
    assert final_compliance_gate(kind,data=original,output_name='Görüş Metni_777.docx',default_name='x.docx',checks=checks,audit_receipt=proof)=='Görüş Metni_777.docx'
    changed=tiny_docx('tip3')
    with pytest.raises(ValueError,match='çıktı değişti'):
        final_compliance_gate(kind,data=changed,output_name='Görüş Metni_777.docx',default_name='x.docx',checks=checks,audit_receipt=proof)
    with pytest.raises(ValueError,match='çıktı değişti'):
        final_compliance_gate(kind,data=original,output_name='Görüş Metni_778.docx',default_name='x.docx',checks=checks,audit_receipt=proof)


def test_no_source_no_proof_even_when_all_flags_true():
    with pytest.raises(ValueError,match='ham kaynak'):
        issue_delivery_evidence('gorus',tiny_docx('gorus'),[],full_checks('gorus'),'Görüş Metni_100.docx')


def test_malicious_edited_receipt_rejected():
    kind='tip3';d=tiny_docx(kind);c=full_checks(kind)
    p=issue_delivery_evidence(kind,d,[b'original bbf'],c,'Ön_Araştırma_Raporu_12.docx')
    p['payload']['facts']['artifact_sha256']='0'*64
    with pytest.raises(ValueError,match='bütünlüğü bozuk'):
        final_compliance_gate(kind,data=d,output_name='Ön_Araştırma_Raporu_12.docx',default_name='x.docx',checks=c,audit_receipt=p)


def test_failed_or_unrun_specialist_check_cannot_issue_proof():
    for kind in FINAL_COMPLIANCE_REQUIRED_CHECKS:
        checks=full_checks(kind); key=next(iter(checks));checks[key]=False
        with pytest.raises(ValueError,match='FAIL/NOT_RUN'):
            issue_delivery_evidence(kind,tiny_docx(kind),[b'original'],checks,'example.docx')


def test_all_user_downloads_have_independent_raw_source_context():
    from pathlib import Path
    app=(Path(__file__).parent/'app.py').read_text(encoding='utf-8')
    worker=(Path(__file__).parent/'job_worker.py').read_text(encoding='utf-8')
    assert app.count('st.download_button(')==1
    assert app.count('compliant_download_button(')==11  # 10 call sites + 1 definition
    assert app.count('audit_sources=')==10
    assert 'issue_delivery_evidence' in app and 'issue_delivery_evidence' in worker
