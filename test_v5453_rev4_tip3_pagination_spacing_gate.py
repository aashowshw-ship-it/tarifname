import io

import pytest
from docx import Document
from docx.shared import Pt

from app_core import (
    ARASTIRMA_TEMPLATE,
    _apply_tip3_page2_layout,
    _validate_research_template_fidelity,
    render_research_docx_smoke_test,
)


def _renderable_template_bytes(*, collapse_gap=False, force_criteria_spill=False):
    doc = Document(str(ARASTIRMA_TEMPLATE))
    _apply_tip3_page2_layout(doc)
    heading_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "2. DEĞERLENDİRME")
    intro_index = next(i for i in range(heading_index + 1, len(doc.paragraphs)) if doc.paragraphs[i].text.strip())
    intro = doc.paragraphs[intro_index]
    intro.text = (
        "Araştırma kapsamında yurtiçi ve yurtdışı patent veritabanlarında taramalar yapılmış, "
        "tespit edilen dokümanlar ekte incelemenize sunulmuştur. Araştırma sonucunda, araştırma konusu "
        "ile teknik yakınlığı en yüksek dokümanlar US11574558B2 (D1) ve US20170296931A1 (D2) "
        "olarak değerlendirilmiştir."
    )
    intro.paragraph_format.line_spacing = 1.5

    if collapse_gap:
        doc.paragraphs[heading_index + 1].paragraph_format.line_spacing = Pt(1)

    if force_criteria_spill:
        keyword_table = doc.tables[0].rows[3].cells[2].tables[0]
        for i, row in enumerate(keyword_table.rows):
            for j, cell in enumerate(row.cells):
                cell.text = (
                    f"very long technical keyword phrase {i}-{j} designed to force the complete criteria "
                    "content beyond the single binding criteria page " * 12
                )

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), doc


def test_evaluation_heading_keeps_one_full_1_5_line_spacer_before_intro():
    doc = Document(str(ARASTIRMA_TEMPLATE))
    _apply_tip3_page2_layout(doc)
    heading_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "2. DEĞERLENDİRME")
    intro_index = next(i for i in range(heading_index + 1, len(doc.paragraphs)) if doc.paragraphs[i].text.strip())
    assert intro_index == heading_index + 2
    assert doc.paragraphs[heading_index + 1].text.strip() == ""
    assert doc.paragraphs[heading_index + 1].paragraph_format.line_spacing == 1.5
    assert doc.paragraphs[intro_index].paragraph_format.line_spacing == 1.5
    assert doc.paragraphs[intro_index + 1].paragraph_format.line_spacing == 1.5
    assert doc.paragraphs[intro_index + 2].paragraph_format.line_spacing == 1.5
    _validate_research_template_fidelity(doc)


def test_docx_gate_rejects_collapsed_evaluation_spacer():
    doc = Document(str(ARASTIRMA_TEMPLATE))
    _apply_tip3_page2_layout(doc)
    heading_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "2. DEĞERLENDİRME")
    doc.paragraphs[heading_index + 1].paragraph_format.line_spacing = Pt(1)
    with pytest.raises(ValueError, match="1,5 satır"):
        _validate_research_template_fidelity(doc)


def test_render_gate_rejects_visible_collapsed_heading_gap():
    data, _ = _renderable_template_bytes(collapse_gap=True)
    with pytest.raises(ValueError, match="dip dibe"):
        render_research_docx_smoke_test(data)


def test_render_gate_rejects_criteria_content_that_spills_beyond_single_page():
    data, _ = _renderable_template_bytes(force_criteria_spill=True)
    with pytest.raises(ValueError, match="kriter sayfası"):
        render_research_docx_smoke_test(data)


def test_render_gate_accepts_binding_two_content_page_start_rhythm():
    data, _ = _renderable_template_bytes()
    assert render_research_docx_smoke_test(data) >= 3
