import io

import pytest
from docx import Document

from app_core import (
    ARASTIRMA_TEMPLATE,
    UploadedAsset,
    _apply_tip3_page2_layout,
    top10_research_prompt,
    validate_research_original_patent_asset,
    validate_research_report_language,
)
from test_v5453_tip3_fail_closed_delivery import _base_report


def _docx_source(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_totalpatent_espaenet_label_is_binding():
    prompt = top10_research_prompt("test invention", "07.09.2026")
    assert '"totalpatent_query":"Totalpatent/Espaenet sorgusu:' in prompt


def test_original_d1_d2_source_must_match_selected_identity():
    good = UploadedAsset(
        "D1.docx",
        _docx_source("US 2022/0238039 A1 ABSTRACT An English abstract section."),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    validate_research_original_patent_asset(good, {"number": "US20220238039A1"}, "D1")

    bad = UploadedAsset(
        "wrong.docx",
        _docx_source("US 2017/0296931 A1 ABSTRACT Another document."),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    with pytest.raises(ValueError, match="doküman kimliği"):
        validate_research_original_patent_asset(bad, {"number": "US20220238039A1"}, "D1")


def test_keyword_table_requires_exactly_ten_terms():
    report = _base_report()
    validate_research_report_language(report)
    report["keywords"] = report["keywords"][:6]
    with pytest.raises(ValueError, match="TAM 10"):
        validate_research_report_language(report)


def test_three_substantive_inventive_step_paragraphs_are_required_for_d1_d2():
    report = _base_report()
    second = dict(report["documents"][0])
    second["label"] = "D2"
    second["number"] = "CN2"
    second["novelty_assessment"] = ["D2 teknik olarak yakın bir karar yapısı açıklamaktadır. Ancak GWP entegrasyonu ve bağlamsal proses seçimi birlikte verilmemektedir. D2 dokümanında GWP entegrasyonu ve bağlamsal proses seçimi ile ilgili bir emareye rastlanmamıştır. Bu kapsamda araştırma konusu buluşun D2 dokümanı varlığında yeni olduğu düşünülmektedir."]
    report["documents"].append(second)
    report["inventive_step_paragraphs"] = [
        "D1 ve D2 birlikte düşünüldüğünde çözüm öngörülebilir görünmektedir. Bu nedenle buluş basamağı yoktur. Kısa değerlendirmedir."
    ]
    with pytest.raises(ValueError, match="TAM 3 güçlü paragraf"):
        validate_research_report_language(report)


def test_evaluation_heading_is_forced_to_new_page_and_intro_stays_with_it():
    doc = Document(str(ARASTIRMA_TEMPLATE))
    _apply_tip3_page2_layout(doc)
    heading_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "2. DEĞERLENDİRME")
    intro_index = next(i for i in range(heading_index + 1, len(doc.paragraphs)) if doc.paragraphs[i].text.strip())
    assert doc.paragraphs[heading_index].paragraph_format.page_break_before is True
    assert doc.paragraphs[intro_index].paragraph_format.page_break_before is not True
