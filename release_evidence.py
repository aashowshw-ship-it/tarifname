"""v5.4.87: Fail-closed, source-bound, independently re-executed delivery audit.

This module is the only issuer of process-local audit receipts. Raw bools returned
by upstream workflows are never accepted as an independent final audit. The audit
reads the actual final OOXML bytes and requires an identified original source.
Semantic source/claim audits remain mandatory upstream; this gate adds a second,
non-skippable verification lane rather than claiming full legal correctness.
"""
from __future__ import annotations

import hashlib
import hmac
import io
import json
import re
import secrets
import zipfile
from collections import Counter
from typing import Any
from docx import Document

_KEY = secrets.token_bytes(32)
_SCHEMA = "delivery-evidence-1"


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _source_bytes(value: Any) -> bytes:
    if isinstance(value, bytes):
        return value
    if isinstance(value, bytearray):
        return bytes(value)
    if isinstance(value, str):
        return value.encode('utf-8')
    if isinstance(value, dict) and isinstance(value.get('data'), (bytes, bytearray)):
        return bytes(value['data'])
    if hasattr(value, 'data') and isinstance(value.data, (bytes, bytearray)):
        return bytes(value.data)
    if hasattr(value, 'getvalue'):
        return bytes(value.getvalue())
    raise ValueError('Teslim kilidi: ham kaynak baytları doğrulanamadı.')


def independent_word_audit(kind: str, data: bytes, sources: list[Any]) -> dict:
    """Actual bytes are reopened. Errors prevent a proof; no AI PASS claims accepted."""
    if not sources:
        raise ValueError('Teslim kilidi: bağımsız kaynak incelemesi için ham kaynak kaydı yok.')
    originals = [_source_bytes(x) for x in sources if x is not None]
    if not originals or any(not x for x in originals):
        raise ValueError('Teslim kilidi: boş veya eksik ham kaynak bulundu.')
    with zipfile.ZipFile(io.BytesIO(bytes(data))) as zf:
        if zf.testzip() is not None:
            raise ValueError('Teslim kilidi: Word ZIP bütünlüğü bozuk.')
        names = zf.namelist()
        if 'word/document.xml' not in names:
            raise ValueError('Teslim kilidi: Word içeriği yok.')
        xml = zf.read('word/document.xml').decode('utf-8')
        if re.search(r'<w:(?:del|ins)(?:\s|>)', xml) and kind == 'figures':
            raise ValueError('Teslim kilidi: şekil belgesinde beklenmeyen revizyon işaretleri var.')
        media = [name for name in names if name.startswith('word/media/') and not name.endswith('/')]
        if kind in {'figures', 'figure_update'} and not media and '<w:drawing' not in xml:
            raise ValueError('Teslim kilidi: şekiller Word belgesinde çizim bulunamadı.')
        if kind == 'tarifname_update' and not re.search(r'<w:(?:ins|del)(?:\s|>)', xml):
            raise ValueError('Teslim kilidi: gerçek Word Track Changes bulunamadı.')
        # Artifacts commonly include the same picture twice intentionally for
        # perspective/overview. Flag them for an explicit figure review upstream;
        # never delete or merge source artwork automatically here.
        repeats = sum(n-1 for n in Counter(_sha(zf.read(name)) for name in media).values() if n > 1)
    doc = Document(io.BytesIO(bytes(data)))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    if not paragraphs:
        raise ValueError('Teslim kilidi: nihai Word belgesinin görünür içeriği yok.')
    full = '\n'.join(paragraphs)
    if kind == 'tarifname':
        for heading in ('TEKNİK ALAN', 'ÖNCEKİ TEKNİK', 'BULUŞUN DETAYLI AÇIKLAMASI', 'İSTEMLER', 'ÖZET'):
            if heading not in full:
                raise ValueError('Teslim kilidi: zorunlu tarifname başlığı eksik: ' + heading)
        claim_text = full.split('İSTEMLER', 1)[-1].split('ÖZET', 1)[0]
        if re.search(r'\b(?:belirlemesidir|sağlamasıdır|oluşturmasıdır)\s*\.', claim_text, re.I):
            raise ValueError('Teslim kilidi: istemde yasak kapanış kalıbı bulundu.')
        # All official template's red/blue explanatory text must remain.
        if not ('Eğer istemler başlığı altında' in claim_text and
                'Buluşu 1. İstemde' in claim_text):
            raise ValueError('Teslim kilidi: istemlerdeki bağlayıcı açıklama metinleri eksik.')
        colors = set()
        for p in doc.paragraphs:
            for run in p.runs:
                if run.text.strip() and run.font.color and run.font.color.rgb:
                    colors.add(str(run.font.color.rgb).upper())
        if not any(c.endswith('FF0000') or c == 'FF0000' for c in colors):
            raise ValueError('Teslim kilidi: kırmızı şablon metni bulunamadı.')
        if not any(c.endswith('0000FF') or c == '0000FF' for c in colors):
            raise ValueError('Teslim kilidi: mavi şablon metni bulunamadı.')
    if kind in ('figures', 'figure_update') and not re.search(r'ŞEKİL\s*\d+', full, re.I):
        raise ValueError('Teslim kilidi: şekil başlığı/numarası bulunamadı.')
    if kind == 'gorus' and not (re.search(r'\b(?:D1|D2|D3)\b', full) or re.search('GÖRÜŞ|RESPONSE', full, re.I)):
        raise ValueError('Teslim kilidi: görüş belgesi içeriği doğrulanamadı.')
    if kind in ('tip3', 'tip3_update') and 'ARAŞTIRMA' not in full.upper():
        raise ValueError('Teslim kilidi: Tip 3 raporu başlığı doğrulanamadı.')
    return {
        'artifact_sha256': _sha(bytes(data)),
        'source_sha256': sorted(_sha(x) for x in originals),
        'paragraph_count': len(paragraphs),
        'media_count': len(media),
        'identical_media_repeats': repeats,
        'auditor': _SCHEMA,
        'kind': kind,
    }


def issue_delivery_evidence(kind: str, data: bytes, sources: list[Any], checks: dict, filename: str) -> dict:
    from rules import FINAL_COMPLIANCE_REQUIRED_CHECKS
    required = FINAL_COMPLIANCE_REQUIRED_CHECKS.get(kind)
    if required is None:
        raise ValueError('Teslim kilidi: tanımsız iş türü.')
    failed = [key for key in required if checks.get(key) is not True]
    if failed:
        raise ValueError('Teslim kilidi: zorunlu kalite kontrolü FAIL/NOT_RUN: ' + ', '.join(failed))
    facts = independent_word_audit(kind, data, sources)
    payload = {'schema': _SCHEMA, 'name': filename, 'type': kind,
               'checks': list(required), 'facts': facts}
    message = json.dumps(payload, sort_keys=True, ensure_ascii=False, separators=(',', ':')).encode()
    return {'payload': payload, 'seal': hmac.new(_KEY, message, hashlib.sha256).hexdigest()}


def verify_delivery_evidence(receipt: dict | None, kind: str, data: bytes, filename: str, checks: dict) -> None:
    if not isinstance(receipt, dict) or not isinstance(receipt.get('payload'), dict):
        raise ValueError('Teslim kilidi: bağımsız denetim makbuzu yok (NOT_RUN).')
    payload = receipt['payload']
    raw = json.dumps(payload, sort_keys=True, ensure_ascii=False, separators=(',', ':')).encode()
    expected = hmac.new(_KEY, raw, hashlib.sha256).hexdigest()
    if not hmac.compare_digest(str(receipt.get('seal', '')), expected):
        raise ValueError('Teslim kilidi: denetim makbuzunun bütünlüğü bozuk.')
    from rules import FINAL_COMPLIANCE_REQUIRED_CHECKS
    if (payload.get('schema') != _SCHEMA or payload.get('type') != kind or
            payload.get('name') != filename or payload.get('checks') != list(FINAL_COMPLIANCE_REQUIRED_CHECKS[kind]) or
            payload.get('facts', {}).get('artifact_sha256') != _sha(bytes(data)) or
            any(checks.get(k) is not True for k in FINAL_COMPLIANCE_REQUIRED_CHECKS[kind])):
        raise ValueError('Teslim kilidi: çıktı değişti veya kanıt başka bir dosya/işlem türüne ait.')
    if not payload.get('facts', {}).get('source_sha256'):
        raise ValueError('Teslim kilidi: kaynak parmak izi yok.')
