from pathlib import Path
from docx import Document
import io
import app_core as app
from rules import APP_VERSION, RULESET_VERSION
from gorus_audit import validate_gorus_template_fidelity

ROOT = Path(__file__).parent

def _opinion():
    return {
        "_output_language": "İngilizce",
        "_recipient_office": "UNITED STATES PATENT AND TRADEMARK OFFICE",
        "application_no": "18/978,599",
        "applicant": "CONRAD MÜHENDİSLİK SANAYİ VE TİCARET LİMİTED ŞİRKETİ",
        "reference": "175179",
        "intro": "In the non-final Office Action, claims 1-13 were rejected under 35 U.S.C. § 103.",
        "amendment_assessment": {"heading":"", "blocks":[]},
        "cited_documents": [],
        "sections": [],
        "combined_assessment": {"heading":"", "paragraphs":[], "groups":[]},
        "conclusion": ["Applicant respectfully requests reconsideration."],
        "signoff": "Respectfully submitted,\nDESTEK PATENT A.Ş.",
    }

def test_version_and_ruleset():
    assert APP_VERSION == "v5.4.73"
    assert RULESET_VERSION == "2026-09-21.v65"

def test_uspto_inference_and_english_shell_passes():
    office = app.infer_gorus_target_office("UNITED STATES PATENT AND TRADEMARK OFFICE\n35 U.S.C. § 103", "Yurtdışı ofis aksiyon")
    assert office == "UNITED STATES PATENT AND TRADEMARK OFFICE"
    op = _opinion()
    data = app.build_gorus_docx(op)
    validate_gorus_template_fidelity(data, ROOT / "Gorus_metni_696809_template.docx", op)
    doc = Document(io.BytesIO(data))
    texts=[p.text.strip() for p in doc.paragraphs]
    assert texts[0] == "UNITED STATES PATENT AND TRADEMARK OFFICE"
    assert texts[1] == "RESPONSE LETTER"
    assert "Dear Examiner," in texts
    assert "Respectfully submitted," in texts
    assert [doc.tables[0].rows[i].cells[0].text.strip() for i in range(3)] == ["Application No.", "Applicant", "Reference"]

def test_foreign_office_unknown_fails_closed():
    try:
        app.infer_gorus_target_office("UNKNOWN FOREIGN OFFICE ACTION", "Yurtdışı ofis aksiyon")
    except ValueError as e:
        assert "hedef-ofis" in str(e)
    else:
        raise AssertionError("unknown foreign office must fail closed")

def test_english_shell_rejects_turkish_template_leak():
    op = _opinion()
    data = app.build_gorus_docx(op)
    doc = Document(io.BytesIO(data))
    doc.paragraphs[3].text = "Sayın Uzman,"
    out=io.BytesIO(); doc.save(out)
    try:
        validate_gorus_template_fidelity(out.getvalue(), ROOT / "Gorus_metni_696809_template.docx", op)
    except ValueError as e:
        assert "hitap" in str(e) or "İngilizce" in str(e)
    else:
        raise AssertionError("Turkish shell leak must fail")
