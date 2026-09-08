from docx import Document

from app_core import ARASTIRMA_TEMPLATE, _apply_tip3_page2_layout, validate_research_report_language
from rules import RULESET_VERSION
from test_v5453_tip3_fail_closed_delivery import _base_report


def test_rev4_ruleset_version():
    assert RULESET_VERSION == "2026-09-08.v46"


def test_evaluation_starts_on_separate_page_not_intro():
    doc = Document(str(ARASTIRMA_TEMPLATE))
    _apply_tip3_page2_layout(doc)
    heading_index = next(i for i,p in enumerate(doc.paragraphs) if p.text.strip() == "2. DEĞERLENDİRME")
    intro_index = next(i for i in range(heading_index+1, len(doc.paragraphs)) if doc.paragraphs[i].text.strip())
    assert doc.paragraphs[heading_index].paragraph_format.page_break_before is True
    assert doc.paragraphs[intro_index].paragraph_format.page_break_before is not True


def test_inventive_step_requires_three_full_paragraphs_not_short_summaries():
    report = _base_report()
    second = dict(report["documents"][0])
    second["label"] = "D2"
    second["number"] = "CN2"
    second["novelty_assessment"] = ["D2 teknik olarak yakındır. D2 dokümanında hedef özellikler ile ilgili bir emareye rastlanmamıştır. Bu kapsamda araştırma konusu buluşun D2 dokümanı varlığında yeni olduğu düşünülmektedir."]
    report["documents"].append(second)
    report["inventive_step_paragraphs"] = [
        "D1 yakındır. Bir fark vardır. Teknik problem vardır. Sonuç öngörülebilir.",
        "D2 tamamlar. Motivasyon vardır. Birleşim mümkündür. Sonuç öngörülebilir.",
        "Kalan özellikler olağandır. Sinerji yoktur. Teknik etki beklenendir. Buluş basamağı yoktur.",
    ]
    try:
        validate_research_report_language(report)
    except ValueError as exc:
        assert "kısa/yüzeysel" in str(exc)
    else:
        raise AssertionError("Kısa buluş basamağı paragrafları kapıdan geçmemelidir")
