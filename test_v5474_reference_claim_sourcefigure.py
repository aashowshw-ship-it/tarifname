from pathlib import Path
import pytest
import app_core as app


def test_reference_display_order_numeric_then_symbolic_then_method_is_separate_block():
    rows = [
        {"number":"K","name":"Arayan"},
        {"number":"160","name":"Sınıf koşullu modül"},
        {"number":"62","name":"Ağırlıklı dal"},
        {"number":"60","name":"Havuzlama başlığı"},
        {"number":"80b","name":"Karar birleştirici"},
        {"number":"80a","name":"Yayın birleştirici"},
        {"number":"O","name":"Operatör"},
        {"number":"61","name":"Dikkat dalı"},
    ]
    ordered = app._ordered_reference_elements(rows)
    assert [r["number"] for r in ordered] == ["60","61","62","80a","80b","160","K","O"]


def test_generic_system_whole_cannot_be_numbered_element():
    draft = {"elements":[{"number":"1","name":"Sistem"}], "system_claim":{}, "dependent_system_claims":[]}
    with pytest.raises(ValueError, match="buluşun bütünü"):
        app._validate_reference_role_rules(draft, "Türkçe")


def test_human_actor_reference_cannot_be_claim_element_but_terminal_can():
    bad = {
        "elements":[{"number":"O","name":"Operatör"}],
        "system_claim":{"preamble":"Teknik bir sistem", "elements":["operatör (O),"], "closing":"içermesidir."},
        "dependent_system_claims":[],
    }
    with pytest.raises(ValueError, match="insan rolü"):
        app._validate_reference_role_rules(bad, "Türkçe")

    good = {
        "elements":[{"number":"10","name":"Operatör terminali"}],
        "system_claim":{"preamble":"Teknik bir sistem", "elements":["operatör terminali (10),"], "closing":"içermesidir."},
        "dependent_system_claims":[],
    }
    app._validate_reference_role_rules(good, "Türkçe")


def test_dependent_claim_type_closure_rejects_generic_structure_and_mismatch():
    draft = {"elements":[{"number":"20","name":"Kanal çözme modülü"}]}
    with pytest.raises(ValueError, match="bir yapıda olmasıdır"):
        app._validate_dependent_system_claim_type_closure([
            "İstem 1’e uygun sistem olup, özelliği; kanal çözme modülünün (20), dar bant veriyi işleyen bir yapıda olmasıdır."
        ], draft, "Türkçe")
    with pytest.raises(ValueError, match="bir modül olmasıdır"):
        app._validate_dependent_system_claim_type_closure([
            "İstem 1’e uygun sistem olup, özelliği; kanal çözme modülünün (20), dar bant veriyi işleyen bir birim olmasıdır."
        ], draft, "Türkçe")
    app._validate_dependent_system_claim_type_closure([
        "İstem 1’e uygun sistem olup, özelliği; kanal çözme modülünün (20), dar bant veriyi işleyen bir modül olmasıdır."
    ], draft, "Türkçe")


def test_additional_independent_method_claims_are_exposed_to_validators():
    draft = {
        "method_claim":{"preamble":"Birinci yöntem için yeterince uzun teknik bir giriş bağlamı sağlayan yöntem", "steps":["birinci işlemin yapılması (1001)"]},
        "dependent_method_claims":["İstem 1’e uygun yöntem olup, özelliği; ek işlem adımını içermesidir."],
        "additional_method_claims":[{
            "claim":{"preamble":"İkinci bağımsız yöntem için yeterince uzun teknik bir giriş bağlamı sağlayan yöntem", "steps":["ikinci işlemin yapılması (1030)"], "closing":"işlem adımlarını içermesidir."},
            "dependent_claims":[]
        }]
    }
    assert len(app._all_method_claim_objects(draft)) == 2
    assert len(app._all_method_dependent_claims(draft)) == 1


def test_text_heavy_source_figure_is_not_automatically_excluded_anymore():
    audit = {
        "final_use":"include",
        "text_heavy_nonreference_numbering":True,
        "nonreference_numeric_or_step_marks":["A1","B2"],
    }
    assert app._figure_audit_requests_exclusion(audit) is False


def test_figure_flow_receives_extra_instruction_and_geometry_fail_closed_fields():
    source = Path(app.__file__).read_text(encoding="utf-8")
    assert "extra_instruction=extra_instruction" in source
    assert 'bool(audit.get("text_cleanup_required"))' in source
    for token in ["arrows_preserved", "box_edges_preserved", "connections_preserved", "only_requested_text_changed"]:
        assert token in source
    assert '"technical_term_clarity_passed"' in source


def test_symbolic_context_refs_are_not_forced_into_system_claims():
    draft = {
        "elements":[
            {"number":"10","name":"Medya çoğaltma birimi"},
            {"number":"K","name":"Arayan"},
            {"number":"O","name":"Operatör"},
            {"number":"R","name":"Ses kaydı"},
        ],
        "system_claim":{
            "preamble":"Telefon kanalı üzerinde sentetik ses tespiti gerçekleştiren bir sistem",
            "elements":["medya çoğaltma birimi (10),"],
            "closing":"içermesidir.",
        },
        "dependent_system_claims":[],
    }
    app._validate_all_elements_covered_in_claims(draft)


def test_numeric_technical_ref_still_must_be_covered_in_system_claims():
    draft = {
        "elements":[
            {"number":"10","name":"Medya çoğaltma birimi"},
            {"number":"20","name":"Kanal çözme modülü"},
        ],
        "system_claim":{
            "preamble":"Telefon kanalı üzerinde sentetik ses tespiti gerçekleştiren bir sistem",
            "elements":["medya çoğaltma birimi (10),"],
            "closing":"içermesidir.",
        },
        "dependent_system_claims":[],
    }
    with pytest.raises(ValueError, match="Kanal çözme modülü"):
        app._validate_all_elements_covered_in_claims(draft)


def test_explicit_source_figure_exclusion_requires_new_fail_closed_reason():
    audit = {
        "final_use":"exclude_only_if_nontechnical_or_unrecoverable",
        "exclusion_reason":"yalnız dekoratif logo olup teknik akış/geometri içermiyor",
        "unique_technical_information":[],
        "spec_coverage_evidence":[],
    }
    assert app._figure_audit_requests_exclusion(audit) is True
    assert app._validate_source_figure_exclusion(audit, {}) is None
    bad = dict(audit, exclusion_reason="kötü")
    assert app._validate_source_figure_exclusion(bad, {}) is not None


def test_figure_audit_prompt_receives_exact_extra_instruction():
    prompt = app._figure_reference_audit_prompt(
        {"elements":[], "method_steps":[]},
        1,
        "Türkçe",
        "Çizimler (algoritmalar) üzerinde yazı olmayacak şekilde düzenlenmeli",
    )
    assert "Çizimler (algoritmalar) üzerinde yazı olmayacak şekilde düzenlenmeli" in prompt


def _clarity_draft_with_question():
    return {
        "system_claim": {
            "preamble": "Telefon kanalında sentetik ses tespiti yapan teknik sistem",
            "elements": ["son N pencere skor değerinin ortalamasını hesaplayan karar birleştirici (80b)"],
            "closing": "içermesidir.",
        },
        "dependent_system_claims": [],
        "method_claim": None,
        "dependent_method_claims": [],
        "additional_method_claims": [],
        "clarity_questions": [{
            "claim_number": "1",
            "anchor_text": "son N pencere skor değerinin ortalamasını",
            "question": "Karar hesabında kullanılan N değeri nasıl ve hangi teknik kritere göre belirlenmektedir?",
            "reason": "N istem içinde tanımlanmamış ve kaynakta kesin belirleme kuralı bulunmamaktadır.",
        }],
        "claim_clarity_review": {
            "review_complete": True,
            "all_issues_resolved_or_queried": True,
            "claim_reviews": [{
                "claim_number": "1",
                "status": "customer_question",
                "issue_type": "undefined_variable",
                "problematic_phrase": "son N pencere",
                "source_quote": "",
                "resolution": "",
            }],
        },
        "coverage_audit": {
            "epo_pct_claim_clarity_reviewed": True,
            "epo_pct_claim_clarity_handled": True,
        },
    }


def test_epo_pct_clarity_unresolved_issue_is_allowed_only_with_customer_question():
    draft = _clarity_draft_with_question()
    app._validate_epo_pct_claim_clarity_review(draft, "Türkçe")
    draft["clarity_questions"] = []
    with pytest.raises(ValueError, match="Word comment sorusu"):
        app._validate_epo_pct_claim_clarity_review(draft, "Türkçe")


def test_epo_pct_clarity_resolved_from_source_requires_basis_and_resolution():
    draft = _clarity_draft_with_question()
    draft["clarity_questions"] = []
    draft["claim_clarity_review"]["claim_reviews"][0].update({
        "status": "resolved_from_source",
        "source_quote": "Kaynakta son pencere sayısının önceden belirlenen bir parametre olduğu açıklanmaktadır.",
        "resolution": "N ifadesi önceden belirlenen sayıda son pencere olarak açıklaştırıldı.",
    })
    app._validate_epo_pct_claim_clarity_review(draft, "Türkçe")
    draft["claim_clarity_review"]["claim_reviews"][0]["source_quote"] = ""
    with pytest.raises(ValueError, match="kaynak alıntısı"):
        app._validate_epo_pct_claim_clarity_review(draft, "Türkçe")


def test_clarity_question_is_word_comment_at_relevant_phrase_with_destek_patent_author():
    import io, zipfile
    from docx import Document
    draft = _clarity_draft_with_question()
    doc = Document()
    doc.add_paragraph("İSTEMLER")
    doc.add_paragraph("son N pencere skor değerinin ortalamasını hesaplayan karar birleştirici (80b)")
    app.add_word_comment_at_phrase(
        doc,
        draft["clarity_questions"][0]["anchor_text"],
        draft["clarity_questions"][0]["question"],
        author="Destek Patent",
        initials="DP",
    )
    bio = io.BytesIO(); doc.save(bio); data = bio.getvalue()
    app._validate_epo_pct_clarity_comments_docx(data, draft)
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        comments = zf.read("word/comments.xml").decode("utf-8")
        document = zf.read("word/document.xml").decode("utf-8")
    assert 'w:author="Destek Patent"' in comments
    assert 'w:initials="DP"' in comments
    assert "Karar hesabında kullanılan N değeri" in comments
    assert "commentRangeStart" in document and "son N pencere" in document


def test_epo_pct_clarity_prompt_says_do_not_block_delivery_and_comment_at_issue():
    prompt = app.tarifname_epo_pct_claim_clarity_prompt(
        "Kaynak teknik metin", "", {"technical_facts": []}, _clarity_draft_with_question(), "Sistem ve yöntem", "Türkçe"
    )
    assert "tarifname üretimini DURDURMA" in prompt
    assert "Word'de tam bu `anchor_text` üzerinde `Destek Patent` yazarlı COMMENT" in prompt
    assert "Yeni teknik bilgi" in prompt


# v5.4.77 — central final compliance gate regressions

def _tiny_docx_bytes(text="x"):
    import io
    from docx import Document
    doc=Document(); doc.add_paragraph(text); bio=io.BytesIO(); doc.save(bio); return bio.getvalue()


def test_final_compliance_gate_decodes_url_name_and_preserves_gorus_spaces():
    from rules import final_compliance_gate
    checks={"raw_sources":True,"quotes":True,"spec_basis_coverage":True,"reference_binding":True,"exact_physical_lines":True,"template":True,"content_flow":True,"render":True,"examiner":True}
    name=final_compliance_gate("gorus", data=_tiny_docx_bytes(), output_name="G%C3%B6r%C3%BC%C5%9F%20Metni_700286.docx", default_name="Görüş Metni.docx", checks=checks)
    assert name == "Görüş Metni_700286.docx"
    assert "%" not in name


def test_final_compliance_gate_missing_single_receipt_is_fail_closed():
    from rules import final_compliance_gate
    import pytest
    checks={"raw_sources":True,"quotes":True,"spec_basis_coverage":True,"reference_binding":True,"exact_physical_lines":True,"template":True,"content_flow":True,"render":True,"examiner":False}
    with pytest.raises(ValueError, match="examiner"):
        final_compliance_gate("gorus", data=_tiny_docx_bytes(), output_name="Görüş Metni_1.docx", default_name="x.docx", checks=checks)


def test_tip3_canonical_name_keeps_underscore_convention():
    from rules import final_compliance_gate
    checks={"delivery":True,"render":True}
    name=final_compliance_gate("tip3", data=_tiny_docx_bytes(), output_name="Ön Araştırma Raporu_181612.docx", default_name="x.docx", checks=checks)
    assert name == "Ön_Araştırma_Raporu_181612.docx"


def test_markup_final_gate_rejects_non_destek_patent_author():
    from rules import final_compliance_gate
    import io, zipfile, pytest
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document(); p=doc.add_paragraph(); ins=OxmlElement('w:ins'); ins.set(qn('w:author'),'Other Author'); r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text='x'; r.append(t); ins.append(r); p._p.append(ins)
    bio=io.BytesIO(); doc.save(bio)
    with pytest.raises(ValueError, match="Destek Patent"):
        final_compliance_gate("tarifname_update", data=bio.getvalue(), output_name="Tarifname_markup.docx", default_name="x.docx", checks={"update_result":True})


def test_ui_has_single_download_entrypoint_only():
    from pathlib import Path
    source=Path(__file__).with_name('app.py').read_text(encoding='utf-8')
    assert source.count('st.download_button(') == 1
    assert 'def compliant_download_button' in source
    assert source.count('compliant_download_button(') >= 8


def test_safe_output_name_gorus_no_longer_turns_space_into_underscore():
    assert app.safe_output_name('Görüş%20Metni_700286.docx','x.docx','gorus') == 'Görüş Metni_700286.docx'
