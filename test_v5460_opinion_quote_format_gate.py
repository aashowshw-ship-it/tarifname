from __future__ import annotations

from docx import Document

import pytest

from gorus_audit import (
    validate_opinion_body_not_accidentally_all_bold,
    validate_opinion_quote_run_formatting,
)
from rules import APP_VERSION, RULESET_VERSION, GORUS_RULES


def _opinion(quote: str = "Birebir teknik tarifname pasajı") -> dict:
    return {
        "application_no": "2026/000001",
        "applicant": "ÖRNEK",
        "reference": "699999",
        "cited_documents": [],
        "sections": [
            {
                "heading": "D1 dokümanı hakkında görüşlerimiz:",
                "blocks": [
                    {"type": "paragraph", "text": "Teknik fark açıklanmaktadır."},
                    {
                        "type": "quote",
                        "text": quote,
                        "lead": "sayfa 9, satır 12-19’da bu durum şu şekilde belirtilmiştir:",
                        "attach_to_previous": True,
                    },
                ],
            }
        ],
        "combined_assessment": {},
    }


def _basis_doc(*, lead_bold=False, quote_bold=True, post_bold=False, inherited_lead=False):
    doc = Document()
    doc.styles["Normal"].font.bold = True
    p = doc.add_paragraph()
    r0 = p.add_run(
        "Teknik fark açıklanmaktadır. "
        "sayfa 9, satır 12-19’da bu durum şu şekilde belirtilmiştir: "
    )
    r0.bold = None if inherited_lead else lead_bold
    r1 = p.add_run("“Birebir teknik tarifname pasajı”")
    r1.bold = quote_bold
    r2 = p.add_run(" Bu açıklama teknik katkının işlevsel sonucunu göstermektedir.")
    r2.bold = post_bold
    return doc


def test_release_version_and_rule_text():
    assert APP_VERSION == "v5.4.70"
    assert RULESET_VERSION == "2026-09-16.v63"
    low = GORUS_RULES.casefold()
    assert "yalnız dış `“...”` tırnakları arasındaki birebir tarifname pasajı" in low
    assert "alıntıdan sonra" in low and "normal" in low
    assert "stilinden" in low and "fail" in low


def test_quote_format_gate_accepts_normal_bold_normal_even_when_normal_style_is_bold():
    validate_opinion_quote_run_formatting(_basis_doc(), _opinion())


def test_quote_format_gate_rejects_bold_lead():
    with pytest.raises(ValueError, match="alıntı öncesindeki"):
        validate_opinion_quote_run_formatting(_basis_doc(lead_bold=True), _opinion())


def test_quote_format_gate_rejects_inherited_bold_lead():
    with pytest.raises(ValueError, match="alıntı öncesindeki"):
        validate_opinion_quote_run_formatting(_basis_doc(inherited_lead=True), _opinion())


def test_quote_format_gate_rejects_normal_quote():
    with pytest.raises(ValueError, match="alıntısı kalın"):
        validate_opinion_quote_run_formatting(_basis_doc(quote_bold=False), _opinion())


def test_quote_format_gate_rejects_bold_post_quote_defence():
    with pytest.raises(ValueError, match="alıntısından sonraki"):
        validate_opinion_quote_run_formatting(_basis_doc(post_bold=True), _opinion())


def test_body_gate_rejects_accidentally_all_bold_substantive_paragraph():
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run(
        "Bu teknik savunma paragrafı, istemdeki teknik farkın çalışma ilişkisini ve teknik katkısını "
        "ayrıntılı biçimde açıklayan yeterince uzun bir gövde paragrafıdır ve başlık değildir."
    )
    r.bold = True
    with pytest.raises(ValueError, match="tamamı kalın"):
        validate_opinion_body_not_accidentally_all_bold(doc, _opinion())


def test_body_gate_allows_declared_bold_section_heading():
    doc = Document()
    p = doc.add_paragraph()
    heading = "D1 dokümanı hakkında görüşlerimiz:"
    r = p.add_run(heading)
    r.bold = True
    validate_opinion_body_not_accidentally_all_bold(doc, _opinion())
